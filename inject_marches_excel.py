"""
Fetch marchés BOAMP Grand Est → filtrare BTP → injectare în Excel.
- Adaugă coloana Marchés_publics după CA_annuel (dacă nu există)
- Actualizează prospecții existenți (count marchés)
- Adaugă prospecți noi din marchés BTP (firme cu ≥2 marchés = relevante)
Rulare: python inject_marches_excel.py
"""
import asyncio, re, json, sys
from collections import defaultdict
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
import aiohttp

XLSX = "outil_prospection_v4.xlsx"
GRAND_EST = ["08","10","51","52","54","55","57","67","68","88"]
DEPT_LABELS = {"08":"Ardennes","10":"Aube","51":"Marne","52":"Haute-Marne",
               "54":"Meurthe-et-Moselle","55":"Meuse","57":"Moselle",
               "67":"Bas-Rhin","68":"Haut-Rhin","88":"Vosges"}

BTP_KEYWORDS = [
    "travaux","construction","rénovation","renovation","maçonnerie","maconnerie",
    "chantier","bâtiment","batiment","gros oeuvre","second oeuvre","électricité",
    "electricite","plomberie","chauffage","menuiserie","charpente","couverture",
    "isolation","étanchéité","peinture","carrelage","revêtement","revetement",
    "démolition","demolition","terrassement","voirie","réseaux","canalisation",
    "assainissement","extension","surélévation","réhabilitation","rehabilitation",
    "restructuration","aménagement","amenagement extérieur","VRD","génie civil",
    "genie civil","ouvrage d'art","station epuration","station d'épuration"
]

EXCLUDE_KEYWORDS = [
    "assurance","fourniture de boissons","distributeur","formation","conseil",
    "audit","informatique","mobilier","véhicule","vehicule","nettoyage","gardien",
    "surveillance","livres","documentation","impression","communication","publicité"
]

TODAY = datetime.now().strftime("%d/%m/%Y")


def is_btp(titre: str) -> bool:
    t = titre.lower()
    if any(ex in t for ex in EXCLUDE_KEYWORDS):
        return False
    return any(kw in t for kw in BTP_KEYWORDS)


async def fetch_boamp_grand_est(jours: int = 90) -> list[dict]:
    connector = aiohttp.TCPConnector(ssl=False)
    url = "https://boamp-datadila.opendatasoft.com/api/explore/v2.1/catalog/datasets/boamp/records"
    since = (datetime.now().__class__.today() - __import__("datetime").timedelta(days=jours)).strftime("%Y-%m-%d")
    all_marches = []

    async with aiohttp.ClientSession(connector=connector) as s:
        for dept in GRAND_EST:
            offset = 0
            while True:
                params = {
                    "where": f'nature_libelle = "Résultat de marché" AND dateparution >= "{since}"',
                    "refine": f"code_departement:{dept}",
                    "limit": 100,
                    "offset": offset,
                    "select": "idweb,dateparution,objet,titulaire,code_departement",
                }
                try:
                    async with s.get(url, params=params, timeout=aiohttp.ClientTimeout(total=30)) as r:
                        if r.status != 200:
                            break
                        d = await r.json()
                    recs = d.get("results", [])
                    for rec in recs:
                        titre = rec.get("objet") or ""
                        if not is_btp(titre):
                            continue
                        titulaires = rec.get("titulaire") or []
                        if not titulaires:
                            continue
                        for t in titulaires:
                            all_marches.append({
                                "marche_id": f"BOAMP-{rec.get('idweb','')}",
                                "titre": titre[:150],
                                "attributaire": t.strip(),
                                "siren": "",
                                "dept": dept,
                                "date_attrib": rec.get("dateparution", ""),
                                "source": "BOAMP",
                                "montant": 0,
                            })
                    if len(recs) < 100:
                        break
                    offset += 100
                except Exception:
                    break

    return all_marches


def count_by_company(marches: list[dict]) -> dict[str, list[dict]]:
    by_name = defaultdict(list)
    for m in marches:
        name = m["attributaire"].upper().strip()
        name = re.sub(r'\s+(SAS|SARL|SA|EURL|SNC|SASU|GIE|SCM|SCP|SCOP)\b', '', name).strip()
        by_name[name].append(m)
    return dict(by_name)


def update_excel(marches_btp: list[dict]) -> None:
    wb = load_workbook(XLSX)
    ws = wb["Prospects"]

    # --- Trouve/insère colonne Marchés_publics après CA_annuel ---
    headers = [c.value for c in ws[1]]
    ca_col_idx = None
    marches_col_idx = None
    for i, h in enumerate(headers, start=1):
        if h == "CA_annuel":
            ca_col_idx = i
        if h == "Marchés_publics":
            marches_col_idx = i

    if marches_col_idx is None:
        insert_at = ca_col_idx + 1  # juste après CA_annuel
        ws.insert_cols(insert_at)
        ws.cell(row=1, column=insert_at, value="Marchés_publics")
        ws.cell(row=1, column=insert_at).font = Font(bold=True)
        marches_col_idx = insert_at
        print(f"Coloana Marchés_publics inserată la col {insert_at}")
    else:
        print(f"Coloana Marchés_publics există deja (col {marches_col_idx})")

    # Re-citim headers după inserare
    headers = [c.value for c in ws[1]]

    # --- Index prospecți existenți (SIREN + Société) ---
    existing_sirens = {}
    existing_names = {}
    siren_col = headers.index("SIREN") + 1 if "SIREN" in headers else None
    soc_col = headers.index("Société") + 1 if "Société" in headers else None

    for row_idx in range(2, ws.max_row + 1):
        siren = ws.cell(row=row_idx, column=siren_col).value if siren_col else None
        soc = ws.cell(row=row_idx, column=soc_col).value if soc_col else None
        if siren:
            existing_sirens[str(siren).strip()] = row_idx
        if soc:
            key = re.sub(r'\s+(SAS|SARL|SA|EURL|SASU)\b', '', str(soc).upper().strip())
            existing_names[key] = row_idx

    # --- Groupare marchés pe companie ---
    by_company = count_by_company(marches_btp)
    print(f"\nCompanii BTP cu marchés atribuite Grand Est: {len(by_company)}")

    updated = 0
    added = 0

    for company_name, company_marches in sorted(by_company.items(),
                                                key=lambda x: -len(x[1])):
        nb = len(company_marches)
        dept = company_marches[0]["dept"]
        last_date = sorted(company_marches, key=lambda m: m["date_attrib"], reverse=True)[0]["date_attrib"]
        marches_detail = "; ".join(set(m["titre"][:60] for m in company_marches[:3]))

        # Cherche dans existants
        row_idx = existing_names.get(company_name)

        if row_idx:
            # Actualise compteur marchés
            ws.cell(row=row_idx, column=marches_col_idx, value=nb)
            updated += 1
        else:
            # Ajoute nouveau prospect si ≥ 2 marchés (filtre qualité)
            if nb < 2:
                continue
            new_row = ws.max_row + 1
            col_map = {h: i+1 for i, h in enumerate(headers)}

            def set_col(header, value):
                if header in col_map:
                    ws.cell(row=new_row, column=col_map[header], value=value)

            set_col("Date_création", TODAY)
            set_col("Société", company_name)
            set_col("Département_code", dept)
            set_col("Département_label", DEPT_LABELS.get(dept, ""))
            set_col("Source", "BOAMP")
            set_col("Statut", "À qualifier")
            set_col("Marchés_publics", nb)
            ws.cell(row=new_row, column=marches_col_idx, value=nb)

            # Fundal albastru deschis (Low — necalificat)
            fill = PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid")
            for col in range(1, len(headers) + 1):
                ws.cell(row=new_row, column=col).fill = fill

            added += 1

    print(f"Prospecți actualizați (marchés count): {updated}")
    print(f"Prospecți noi adăugați (≥2 marchés BTP): {added}")

    # --- Feuille Marchés_BOAMP (rezumat complet) ---
    if "Marchés_BOAMP" in wb.sheetnames:
        del wb["Marchés_BOAMP"]
    ws2 = wb.create_sheet("Marchés_BOAMP")
    ws2.append(["Attributaire", "Dept", "Titre", "Date", "Source"])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    ws2["C1"].font = Font(bold=True)
    ws2["D1"].font = Font(bold=True)
    ws2["E1"].font = Font(bold=True)
    for m in sorted(marches_btp, key=lambda x: x["date_attrib"], reverse=True):
        ws2.append([m["attributaire"], m["dept"], m["titre"][:120], m["date_attrib"], m["source"]])

    # Largimi coloane
    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["C"].width = 70
    ws2.column_dimensions["D"].width = 15

    wb.save(XLSX)
    print(f"\nExcel salvat: {XLSX}")
    print(f"Foaie Marchés_BOAMP: {len(marches_btp)} marchés BTP Grand Est")


async def main():
    print("Fetching marchés BOAMP Grand Est (90 zile)...")
    marches = await fetch_boamp_grand_est(jours=90)
    print(f"Total marchés BTP filtrate: {len(marches)}")
    update_excel(marches)


if __name__ == "__main__":
    asyncio.run(main())
