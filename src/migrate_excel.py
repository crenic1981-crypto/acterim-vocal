"""Migrator one-shot: Excel outil_prospection v4 → Google Sheets v3.0.
Rulare: python -m src.migrate_excel <path/to/xlsx>
"""
from __future__ import annotations
import asyncio, sys
from pathlib import Path
from openpyxl import load_workbook
from .google_sheets import SheetsClient, PROSPECTS_HEADERS, fmt_date

PLACEHOLDER = {"← NOUVEAU", "← NOU", "← NOUVE", "← N", None, ""}

DEPT_LABELS = {"08":"Ardennes","10":"Aube","51":"Marne","52":"Haute-Marne",
               "54":"Meurthe-et-Moselle","55":"Meuse","57":"Moselle",
               "67":"Bas-Rhin","68":"Haut-Rhin","88":"Vosges"}

EXCLUS_TOURNEE = {"52","55","10","08"}


def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s in PLACEHOLDER or s.startswith("←"):
        return ""
    return s


def map_prospect_row(r: tuple) -> list:
    """Excel 38 col → Sheets 43 col (PROSPECTS_HEADERS)."""
    # Indici Excel (0-based)
    date_creation = clean(r[0])
    siren         = clean(r[1])
    societe       = clean(r[2])
    adresse       = clean(r[3])
    cp            = clean(r[4])
    ville         = clean(r[5])
    dept          = clean(r[6])
    dept_label    = clean(r[7]) or DEPT_LABELS.get(dept, "")
    zone_tournee  = clean(r[8]) or ("" if dept in EXCLUS_TOURNEE else "tournée_ok")
    metier        = clean(r[9])
    effectif      = clean(r[10])
    ca            = clean(r[11])
    naf           = clean(r[12])
    forme         = clean(r[13])
    dirigeant     = clean(r[14])
    fonction      = clean(r[15]) or "Gérant"
    tel           = clean(r[16])
    email         = clean(r[17])
    c2_nom        = clean(r[18])
    c2_fct        = clean(r[19])
    c2_tel        = clean(r[20])
    source        = clean(r[21]) or "Excel"
    statut        = clean(r[22]) or "À qualifier"
    priorite      = clean(r[23]) or "Moyenne"
    garantie      = clean(r[24])
    montant_g     = clean(r[25])
    score         = clean(r[26]) or 0
    date_modif    = clean(r[27])
    # 28,29 dates garantie — neutilizate direct
    besoin        = clean(r[30])
    accroche      = clean(r[31])
    lat           = clean(r[32])
    lon           = clean(r[33])
    site          = clean(r[34])
    linkedin      = clean(r[35])
    derniere      = clean(r[36])
    # tier
    try:
        sc = int(float(score))
    except (ValueError, TypeError):
        sc = 0
    tier = "High" if sc >= 71 else ("Medium" if sc >= 31 else "Low")

    return [
        siren, societe, adresse, cp, ville, dept, dept_label, zone_tournee,
        metier, naf, forme, effectif, ca, date_creation,
        dirigeant, fonction, tel, email, "false",
        c2_nom, c2_fct, c2_tel,
        site, linkedin,
        sc, tier, 0, source,
        statut, priorite,
        garantie, montant_g,
        fmt_date(), derniere, date_modif,
        besoin, accroche,
        lat, lon,
        "", "", "", "",
    ]


async def migrate(path: str):
    wb = load_workbook(path, data_only=True)
    sheets = SheetsClient()
    await sheets.ensure_structure()

    # --- Prospects ---
    ws = wb["Prospects"]
    existing_sirens = {p["SIREN"] for p in await sheets.all_prospects()}
    rows_added = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not clean(r[1]):  # SIREN absent
            continue
        siren = clean(r[1])
        if siren in existing_sirens:
            continue
        row = map_prospect_row(r)
        await sheets.append("Prospecți", row)
        existing_sirens.add(siren)
        rows_added += 1
    print(f"Prospecți migrați: {rows_added}")

    # --- Scripts ---
    if "Scripts" in wb.sheetnames:
        ws = wb["Scripts"]
        n = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not clean(r[0]):
                continue
            await sheets.append("Scripts", [clean(v) for v in r[:8]])
            n += 1
        print(f"Scripts migrați: {n}")

    # --- Actions & relances ---
    if "Actions & relances" in wb.sheetnames:
        ws = wb["Actions & relances"]
        n = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not clean(r[0]):
                continue
            await sheets.append("Actions", [clean(v) for v in r[:18]])
            n += 1
        print(f"Acțiuni migrate: {n}")

    # --- Tournées terrain ---
    if "Tournées terrain" in wb.sheetnames:
        ws = wb["Tournées terrain"]
        n = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not clean(r[0]):
                continue
            await sheets.append("Tournées", [clean(v) for v in r[:21]])
            n += 1
        print(f"Tournées migrate: {n}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python -m src.migrate_excel <xlsx>")
        sys.exit(1)
    asyncio.run(migrate(sys.argv[1]))


if __name__ == "__main__":
    main()
